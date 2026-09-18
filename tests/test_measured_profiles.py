"""Measured-profile contracts with invented fixture data; no customer measurements in Git."""
from __future__ import annotations

import json
from dataclasses import replace

import openpyxl
import pytest
from click.testing import CliRunner

from solace_autoscale.capacity.benchmarks import (
    BenchmarkPoint,
    BenchmarkSheet,
    BenchmarkWorkbook,
    import_directory,
    import_workbook,
)
from solace_autoscale.capacity.model import lookup
from solace_autoscale.capacity.profile_model import OutsideBenchmark, compile_profile
from solace_autoscale.capacity.schema import CapacityModel
from solace_autoscale.cli import main
from solace_autoscale.config import Config
from solace_autoscale.decision.engine import DecisionRequest, decide
from solace_autoscale.decision.types import Action, MetricSample, ShardInput


@pytest.fixture
def book():
    points = []
    for scenario, factor in [('direct', 2), ('streaming', 1), ('unspooling', .8), ('replay', .6), ('tracing', .5)]:
        for fanout, rate in [(1, 1000), (2, 800), (5, 400)]:
            for size in [1000, 2000]:
                ingress = rate * factor * 1000 / size
                points.append(BenchmarkPoint(scenario=scenario, fanout=fanout, msg_size_bytes=size,
                                             ingress_msg_rate=ingress, egress_msg_rate=ingress * fanout,
                                             ingress_cell=f'C{len(points)+1}', egress_cell=f'K{len(points)+1}'))
    return BenchmarkWorkbook(source_filename='Solace-Cloud-AWS-10.1.2.3.xlsx', source_sha256='a'*64,
                             provider='aws', broker_version='10.1.2.3', sheets=[BenchmarkSheet(
                                 sheet='sc-1k', service_class='enterprise-1k', metadata={'broker mode':'HA'},
                                 observations=points)])


@pytest.fixture
def limits():
    return {'enterprise-1k': {'service_class_id': 'ENTERPRISE_1K_HIGHAVAILABILITY',
                              'connections_max': 1000, 'spool_bytes_max': 1000000000}}


@pytest.fixture
def model(book, limits):
    return compile_profile(book, limits, compiled_at='2026-01-01T00:00:00Z', limits_source='invented test limits')


def test_exact_directions_and_fanout(model):
    p = lookup(model, 'enterprise-1k', 1000, 'guaranteed', fanout=2, scenario='streaming')
    assert p.msg_rate == 800
    assert p.ingress_byte_rate == 800000
    assert p.egress_byte_rate == 1600000
    assert not p.interpolated
    assert p.source_cells


def test_scenarios_are_not_blended(model):
    rates = [lookup(model, 'enterprise-1k', 1000, 'guaranteed', scenario=s).msg_rate
             for s in ['streaming', 'unspooling', 'replay', 'tracing', 'worst']]
    assert rates == [1000, 800, 600, 500, 500]


def test_intermediate_size_respects_both_neighbor_limits(model):
    p = lookup(model, 'enterprise-1k', 1500, 'guaranteed', scenario='streaming')
    assert p.msg_rate == 500  # cannot invent 750 msg/s between measured points
    assert p.ingress_byte_rate <= 1000000
    assert p.interpolated


def test_fractional_fanout_uses_both_neighbors(model):
    p = lookup(model, 'enterprise-1k', 1000, 'guaranteed', fanout=1.5, scenario='streaming')
    assert p.msg_rate == pytest.approx(1000 / 1.5)
    assert p.interpolated


@pytest.mark.parametrize('size,fanout', [(999,1),(2001,1),(1000,6),(1000,.5),(float('nan'),1)])
def test_no_extrapolation(model, size, fanout):
    with pytest.raises(OutsideBenchmark):
        lookup(model, 'enterprise-1k', size, 'guaranteed', fanout=fanout)


def test_missing_scenario_is_not_silently_omitted(book, limits):
    book.sheets[0].observations = [p for p in book.sheets[0].observations if p.scenario != 'tracing']
    m = compile_profile(book, limits, compiled_at='x', limits_source='test')
    with pytest.raises(OutsideBenchmark):
        lookup(m, 'enterprise-1k', 1000, 'guaranteed')
    assert lookup(m, 'enterprise-1k', 1000, 'guaranteed', scenario='replay').msg_rate == 600


def test_model_identity_covers_limits_source_and_measurements(book, limits, model):
    same = compile_profile(book, limits, compiled_at='other time', limits_source='invented test limits')
    assert same.model_version == model.model_version
    limits['enterprise-1k']['connections_max'] = 900
    changed = compile_profile(book, limits, compiled_at='x', limits_source='invented test limits')
    assert changed.model_version != model.model_version
    assert changed.provenance.limits_sha256 != model.provenance.limits_sha256


def test_ha_benchmark_rejects_standalone_limits(book, limits):
    limits['enterprise-1k']['service_class_id'] = 'ENTERPRISE_1K_STANDALONE'
    with pytest.raises(ValueError, match='HA benchmarks'):
        compile_profile(book, limits, compiled_at='x', limits_source='test')


def test_schema_two_cannot_fall_back_to_legacy_curve(model):
    doc = model.model_dump(by_alias=True)
    doc['service_classes']['enterprise-1k']['benchmark'] = None
    with pytest.raises(ValueError, match='benchmark observations'):
        CapacityModel.model_validate(doc)
    doc['schema_version'] = '999'
    with pytest.raises(ValueError, match='unsupported capacity schema'):
        CapacityModel.model_validate(doc)


def _request(model, **changes):
    cfg = Config.model_validate({'fleet': {'service_class':'enterprise-1k'},
                                 'metrics': {'scrape_interval':10},
                                 'capacity': {'scenario':'streaming'},
                                 'policy': {'scale_up_window':30, 'headroom':{'mode':'fixed'}}})
    s = MetricSample(0, 600, 600, 600000, 600000, 1000, 1, 0, 1)
    samples = [replace(s, timestamp=t, **changes) for t in [0,10,20,30]]
    return DecisionRequest(cfg, model, ShardInput('orders', samples), now=30)


def test_one_to_one_bytes_not_double_counted(model):
    d = decide(_request(model))
    assert d.action == Action.hold
    assert d.axes['bytes'].demand_ratio == pytest.approx(.6)
    assert d.recommended_brokers == 1


def test_fanout_pressure_changes_recommendation(model):
    d = decide(_request(model, egress_msg_rate=3000, egress_byte_rate=3000000))
    assert d.action == Action.scale_up
    assert d.recommended_brokers == 2


def test_unsupported_recent_sample_refuses_instead_of_averaging_away(model):
    req = _request(model)
    req.shard.samples[0] = replace(req.shard.samples[0], avg_msg_size=500)
    assert decide(req).action == Action.no_decision


def test_unsupported_old_down_window_holds(model):
    req = _request(model, ingress_msg_rate=1, egress_msg_rate=1, ingress_byte_rate=1000,
                   egress_byte_rate=1000, current_brokers=3)
    req.config.billing.model = 'elastic'
    req.config.policy.scale_down_window = 60
    req = replace(req, now=130)
    req.shard.samples[:] = [replace(s, timestamp=s.timestamp+100) for s in req.shard.samples]
    old = [replace(req.shard.samples[0], timestamp=t, avg_msg_size=999) for t in [70,80,90]]
    req.shard.samples[:0] = old
    d = decide(req)
    assert d.action == Action.hold
    assert 'outside measured' in d.reason


def workbook(tmp_path):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'sc-1k'
    s['B6']='Cloud Service Provider'; s['C6']='Azure'
    s['B24']='Broker Mode'; s['C24']='HA'
    s['B30']='Guaranteed Streaming'
    s['B32']='Ingress Msgs/Sec'; s['J32']='Egress Msgs/Sec'
    s['B33']='Fanout'; s['C33']=1000; s['J33']='Fanout'; s['K33']=1000
    s['B34']=1; s['C34']='12.3K'; s['J34']=1; s['K34']=12200
    p=tmp_path/'Solace-Cloud-Perf-Azure-10.1.2.3.xlsx'
    wb.save(p)
    return p


def test_import_normalizes_compact_rate_and_retains_cells(tmp_path):
    p=workbook(tmp_path)
    b=import_workbook(p)
    point=b.sheets[0].observations[0]
    assert point.ingress_msg_rate == 12300
    assert point.egress_msg_rate == 12200
    assert (point.ingress_cell,point.egress_cell)==('C34','K34')
    assert b.provider == 'azure'


@pytest.mark.parametrize('cell,value', [('C34',None),('C34',0),('C34','=1+1'),('C34','N/A'),
                                      ('B34',1.2),('J34',2),('K33',2000),('J32','Ingress'),('C6','AWS')])
def test_invalid_import_fails_with_context(tmp_path, cell, value):
    p=workbook(tmp_path)
    w=openpyxl.load_workbook(p); w.active[cell]=value; w.save(p); w.close()
    with pytest.raises(ValueError):
        import_workbook(p)


def test_import_validates_every_workbook_before_writing(tmp_path):
    p=workbook(tmp_path)
    bad=tmp_path/'Solace-Cloud-GCP-10.1.2.3.xlsx'
    bad.write_bytes(p.read_bytes())
    out=tmp_path/'out'
    with pytest.raises(ValueError):
        import_directory(tmp_path, out)
    assert not out.exists()


def test_plan_cli_reports_ceiling_and_sources(tmp_path, model):
    path=tmp_path/'model.json'; path.write_text(model.model_dump_json(by_alias=True))
    result=CliRunner().invoke(main,['plan','--model',str(path),'--service-class','enterprise-1k',
                                   '--message-size','1000','--messages','1000','--max-brokers','1'])
    assert result.exit_code == 0, result.output
    data=json.loads(result.output)
    assert data['required_brokers']==3
    assert not data['within_configured_ceiling']
    assert data['source_cells']
    invalid=CliRunner().invoke(main,['plan','--model',str(path),'--service-class','enterprise-1k',
                                    '--message-size','1000','--messages','nan'])
    assert invalid.exit_code != 0


def test_window_average_cannot_hide_slower_large_message_capacity(book, limits):
    sheet = book.sheets[0]
    for p in list(sheet.observations):
        if p.msg_size_bytes == 2000:
            sheet.observations.append(p.model_copy(update={
                'msg_size_bytes': 4000, 'ingress_msg_rate': p.ingress_msg_rate / 2,
                'egress_msg_rate': p.egress_msg_rate / 2,
            }))
    measured = compile_profile(book, limits, compiled_at='test', limits_source='invented limits')
    cfg = Config.model_validate({
        'fleet': {'service_class': 'enterprise-1k'},
        'capacity': {'scenario': 'streaming'},
        'metrics': {'scrape_interval': 10}, 'policy': {'scale_up_window': 30},
        'workload': {'delivery': 'guaranteed'},
    })
    samples = [MetricSample(t, 600, 600, 600 * size, 600 * size, size, 1, 0, 1)
               for t, size in [(0, 1000), (10, 1000), (20, 1000), (30, 4000)]]
    result = decide(DecisionRequest(cfg, measured, ShardInput('orders', samples), now=30))
    assert result.axes['messages'].demand_ratio == pytest.approx(600 / 250)
